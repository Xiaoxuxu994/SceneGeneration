#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# python scripts/summarize_metrics.py --out abc.xlsx

import re
import argparse
from pathlib import Path
from typing import Dict, Tuple, Optional, Set

import numpy as np
import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# 1. 实验配置：主要改这里
# ============================================================

MODEL = "reliev3r"
ENCODER_TYPE = "dinov3"  # "dinov2"
STEPS = [5000, 10000, 15000, 20000, 25000, 30000, 35000, 40000, 45000, 50000]
EXP_DISPLAY_NAMES = {
    "0625_5Bv2_v10-1_stage1": "stage 1",
}
DST_EXP_NAMES = list(EXP_DISPLAY_NAMES.keys())

# summary 输出顺序
DATASETS = [
    "dl3dv",
    "ase",
    "scannet",
    "scannetpp",
    "re10k",
    "nrgbd",
    "dtu",
    "7scenes",
    "eth3d",
    "kitti",
    "hiroom",
    "sekai",
    "bonn"
]

# 每个 dataset 只读取这个 setting 目录
if ENCODER_TYPE == "dinov2":
    DATASET_SETTING_MAP = {
        "7scenes": "test_stride40_maxall_len8_numseq10_518x392",
        "ase": "test_rand_maxall_len8_numseq10_518x518",
        "bonn": None,
        "dl3dv": "test_stride4_maxall_len8_numseq100_518x294",
        "dtu": "test_stride1_maxall_len8_numseq10_518x392",
        "eth3d": "test_stride1_maxall_len8_numseq10_518x350",
        "hiroom": None,
        "kitti": "test_stride2_maxall_len8_numseq10_518x154",
        "nrgbd": "test_stride100_maxall_len8_numseq10_518x392",
        "re10k": "test_stride4_maxall_len8_numseq100_518x294",
        "scannet": "test_stride25_maxall_len8_numseq10_518x392",
        "scannetpp": "test_stride25_maxall_len8_numseq10_518x350",
        "sekai": None,
        # "sintel": "test_stride5_maxall_len8_numseq14_518x224"
    }
elif ENCODER_TYPE == "dinov3":
    DATASET_SETTING_MAP = {
        "7scenes": "test_stride40_maxall_len8_numseq10_512x384",
        "ase": "test_rand_maxall_len8_numseq10_512x512",
        "bonn": "test_stride10_maxall_len8_numseq10_512x384",
        "dl3dv": "test_stride4_maxall_len8_numseq100_512x288",
        "dtu": "test_stride1_maxall_len8_numseq10_512x384",
        "eth3d": "test_stride1_maxall_len8_numseq10_512x336",
        "hiroom": "test_stride1_maxall_len8_numseq30_512x512",
        "kitti": "test_stride2_maxall_len8_numseq10_512x144",
        "nrgbd": "test_stride100_maxall_len8_numseq10_512x384",
        "re10k": "test_stride4_maxall_len8_numseq100_512x288",
        "scannet": "test_stride25_maxall_len8_numseq10_512x384",
        "scannetpp": "test_stride25_maxall_len8_numseq10_512x336",
        "sekai": "test_stride4_maxall_len8_numseq100_512x288",
        # "sintel": "test_stride5_maxall_len8_numseq14_518x224"
    }

DATASET_DISPLAY = {
    "dl3dv": "DL3DV",
    "ase": "ASE",
    "bonn": "Bonn",
    "hiroom": "HiRoom",
    "scannet": "ScanNet",
    "scannetpp": "ScanNet++",
    "sekai": "Sekai",
    "re10k": "RE10K",
    "nrgbd": "NRGBD",
    "dtu": "DTU",
    "7scenes": "7Scenes",
    "eth3d": "ETH3D",
    "kitti": "KITTI",
    "sintel": "Sintel",
}

IN_DOMAIN = {"dl3dv", "ase", "scannet", "scannetpp"}
OUT_DOMAIN = {"re10k", "nrgbd", "dtu", "7scenes", "eth3d", "kitti", "hiroom", "sekai", "bonn"}

DOMAIN_ORDER = {
    "In-domain": 0,
    "Out-of-domain": 1,
    "Other": 2,
}


# ============================================================
# 2. 指标配置
# ============================================================

METRIC_SPECS = [
    {
        "task": "pose",
        "metric": "RRA@30↑",
        "stat": "",
        "key": "rra30",
        "aliases": ["RRA@30↑", "RRA@30", "rra30"],
    },
    {
        "task": "pose",
        "metric": "RTA@30↑",
        "stat": "",
        "key": "rta30",
        "aliases": ["RTA@30↑", "RTA@30", "rta30"],
    },
    {
        "task": "pose",
        "metric": "AUC@30↑",
        "stat": "",
        "key": "auc30",
        "aliases": ["AUC@30↑", "AUC@30", "auc30"],
    },
    {
        "task": "pose",
        "metric": "ATE↓",
        "stat": "",
        "key": "ate",
        "aliases": ["ATE↓", "ATE", "ate"],
    },
    {
        "task": "pose",
        "metric": "RPE trans↓",
        "stat": "",
        "key": "rpe_trans",
        "aliases": ["RPE trans↓", "RPE trans", "rpe trans", "rpe_trans"],
    },
    {
        "task": "pose",
        "metric": "RPE rot↓",
        "stat": "",
        "key": "rpe_rot",
        "aliases": ["RPE rot↓", "RPE rot", "rpe rot", "rpe_rot"],
    },
    {
        "task": "pointmap",
        "metric": "Acc.↓",
        "stat": "Mean",
        "key": "acc_mean",
        "aliases": ["Acc.↓", "Acc.", "Acc", "accuracy"],
    },
    {
        "task": "pointmap",
        "metric": "Acc.↓",
        "stat": "Med.",
        "key": "acc_med",
        "aliases": ["Acc.↓", "Acc.", "Acc", "accuracy"],
    },
    {
        "task": "pointmap",
        "metric": "Comp.↓",
        "stat": "Mean",
        "key": "comp_mean",
        "aliases": ["Comp.↓", "Comp.", "Comp", "completion"],
    },
    {
        "task": "pointmap",
        "metric": "Comp.↓",
        "stat": "Med.",
        "key": "comp_med",
        "aliases": ["Comp.↓", "Comp.", "Comp", "completion"],
    },
    {
        "task": "pointmap",
        "metric": "N.C.↑",
        "stat": "Mean",
        "key": "nc_mean",
        "aliases": ["N.C.↑", "N.C.", "NC", "normal consistency"],
    },
    {
        "task": "pointmap",
        "metric": "N.C.↑",
        "stat": "Med.",
        "key": "nc_med",
        "aliases": ["N.C.↑", "N.C.", "NC", "normal consistency"],
    },
    {
        "task": "depth",
        "metric": "Abs Rel↓",
        "stat": "",
        "key": "abs_rel",
        "aliases": ["Abs Rel↓", "Abs Rel", "AbsRel", "abs_rel"],
    },
    {
        "task": "depth",
        "metric": "δ < 1.25↑",
        "stat": "",
        "key": "delta1",
        "aliases": ["δ < 1.25↑", "δ < 1.25", "delta < 1.25", "delta1", "a1"],
    },
]

METRIC_KEY_TO_SPEC = {m["key"]: m for m in METRIC_SPECS}
METRIC_KEY_ORDER = {m["key"]: i for i, m in enumerate(METRIC_SPECS)}


# ============================================================
# 3. 通用工具
# ============================================================

def norm_text(x) -> str:
    if pd.isna(x):
        return ""

    s = str(x).strip().lower()
    s = s.replace("δ", "delta")
    s = s.replace("↑", "")
    s = s.replace("↓", "")
    s = s.replace("<", "")
    s = s.replace(">", "")
    s = re.sub(r"[^a-z0-9]+", "", s)
    return s


def step_to_str(step: int) -> str:
    if step == 0:
        return "0"
    if step % 1000 == 0:
        return f"{step // 1000}k"
    return str(step)


def get_domain(dataset: str) -> str:
    d = dataset.lower()
    if d in IN_DOMAIN:
        return "In-domain"
    if d in OUT_DOMAIN:
        return "Out-of-domain"
    return "Other"


def parse_setting(setting_name: str) -> str:
    s = setting_name
    parts = []

    stride = re.search(r"stride(\d+)", s)
    rand_max = re.search(r"rand_max(\d+)", s)
    is_randall = "randall" in s or "rand_maxall" in s
    is_rand = "rand" in s

    if stride:
        parts.append(f"stride={stride.group(1)}")
    elif rand_max:
        parts.append(f"random={rand_max.group(1)}")
    elif is_randall:
        parts.append("random=all")
    elif is_rand:
        parts.append("random")

    len_match = re.search(r"len(\d+)", s)
    max_match = re.search(r"(?:^|_)max(\d+)(?:_|$)", s)

    if len_match:
        parts.append(f"seq_len={len_match.group(1)}")
    elif max_match:
        parts.append(f"seq_len={max_match.group(1)}")

    numseq = re.search(r"numseq(\d+)", s)
    if numseq:
        parts.append(f"num_seq={numseq.group(1)}")

    return ", ".join(parts) if parts else setting_name


def expected_xlsx_path(root: Path, dataset: str, exp_name: str, step: int) -> Path:
    setting_dir = DATASET_SETTING_MAP[dataset]
    return root / dataset / setting_dir / f"metrics_{MODEL}_{exp_name}_step{step}.xlsx"


# ============================================================
# 4. 读取单个 xlsx：只读 summary sheet
# ============================================================

def find_summary_sheet_name(xls: pd.ExcelFile) -> Optional[str]:
    for s in xls.sheet_names:
        if str(s).strip().lower() == "summary":
            return s

    for s in xls.sheet_names:
        if "summary" in str(s).strip().lower():
            return s

    return None


def extract_metrics_from_xlsx(xlsx_path: Path) -> Optional[Tuple[Dict[str, float], Set[str]]]:
    """
    只读取 summary sheet。

    期望 summary sheet:
    A: dataset
    B: task
    C: metric
    D: stat
    E: value

    返回:
      metrics: {metric_key: value}
      present_keys: 当前 sheet 里真实出现过的 metric_key
    """
    metrics = {spec["key"]: np.nan for spec in METRIC_SPECS}
    present_keys: Set[str] = set()

    try:
        xls = pd.ExcelFile(xlsx_path)
    except Exception as e:
        print(f"[WARN] failed to open xlsx, keep blank: {xlsx_path}, error={e}")
        return None

    summary_sheet = find_summary_sheet_name(xls)

    if summary_sheet is None:
        print(f"[WARN] no summary sheet, keep blank: {xlsx_path}")
        print(f"       available sheets: {xls.sheet_names}")
        return None

    try:
        df = pd.read_excel(
            xlsx_path,
            sheet_name=summary_sheet,
            header=None,
            usecols=[0, 1, 2, 3, 4],
        )
    except Exception as e:
        print(f"[WARN] failed to read summary sheet, keep blank: {xlsx_path}, error={e}")
        return None

    df = df.dropna(how="all")

    if df.empty:
        print(f"[WARN] empty summary sheet, keep blank: {xlsx_path}")
        return None

    df.columns = ["dataset_in_sheet", "task", "metric", "stat", "value"]

    # 处理合并单元格读出来的 NaN
    df["dataset_in_sheet"] = df["dataset_in_sheet"].ffill()
    df["task"] = df["task"].ffill()
    df["metric"] = df["metric"].ffill()
    df["stat"] = df["stat"].fillna("")
    df["value"] = pd.to_numeric(df["value"], errors="coerce")

    # 保留 value 可转数字的行
    df = df[df["value"].notna()].copy()

    if df.empty:
        print(f"[WARN] no numeric metric values, keep blank: {xlsx_path}")
        return None

    for spec in METRIC_SPECS:
        key = spec["key"]
        target_task = norm_text(spec["task"])
        target_stat = norm_text(spec["stat"])
        target_aliases = [norm_text(a) for a in spec["aliases"]]

        matched_values = []

        for _, row in df.iterrows():
            row_task = norm_text(row["task"])
            row_metric = norm_text(row["metric"])
            row_stat = norm_text(row["stat"])

            if target_task and row_task != target_task:
                continue

            metric_hit = False
            for alias in target_aliases:
                if not alias:
                    continue
                if row_metric == alias or alias in row_metric or row_metric in alias:
                    metric_hit = True
                    break

            if not metric_hit:
                continue

            if target_stat:
                if row_stat != target_stat and target_stat not in row_stat:
                    continue

            matched_values.append(float(row["value"]))

        if matched_values:
            metrics[key] = matched_values[-1]
            present_keys.add(key)

    return metrics, present_keys


# ============================================================
# 5. 收集所有结果
# ============================================================

def collect_records(root: Path) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """
    只读指定 dataset / setting / exp / step。

    xlsx 不存在：
      - 不报错
      - 对应结果留空
      - 继续下一个
    """
    file_cache = {}
    dataset_present_metric_keys = {dataset: set() for dataset in DATASETS}
    selected_file_records = []

    expected_num = len(DATASETS) * len(DST_EXP_NAMES) * len(STEPS)
    missing_files = []
    failed_files = []

    for dataset in DATASETS:
        setting_dir = DATASET_SETTING_MAP[dataset]
        metric_dir = root / dataset / setting_dir

        if not metric_dir.exists():
            print(f"[WARN] metric dir not found, keep blank: {metric_dir}")

        for exp_name in DST_EXP_NAMES:
            for step in STEPS:
                xlsx_path = expected_xlsx_path(root, dataset, exp_name, step)

                exists = xlsx_path.exists()
                parsed_ok = False

                metrics = {spec["key"]: np.nan for spec in METRIC_SPECS}
                present_keys: Set[str] = set()

                if not exists:
                    missing_files.append(str(xlsx_path))
                else:
                    parsed = extract_metrics_from_xlsx(xlsx_path)
                    if parsed is None:
                        failed_files.append(str(xlsx_path))
                    else:
                        metrics, present_keys = parsed
                        parsed_ok = True
                        dataset_present_metric_keys[dataset].update(present_keys)

                file_cache[(dataset, exp_name, step)] = {
                    "metrics": metrics,
                    "exists": exists,
                    "parsed_ok": parsed_ok,
                    "xlsx": str(xlsx_path),
                }

                selected_file_records.append({
                    "dataset": dataset,
                    "setting_dir": setting_dir,
                    "setting": parse_setting(setting_dir),
                    "exp": exp_name,
                    "step": step,
                    "exists": exists,
                    "parsed_ok": parsed_ok,
                    "xlsx": str(xlsx_path),
                })

    records = []

    for dataset in DATASETS:
        setting_dir = DATASET_SETTING_MAP[dataset]
        setting = parse_setting(setting_dir)

        # 若该 dataset 至少有一个 xlsx 成功解析，则只输出真实出现过的 metric。
        # 若该 dataset 全部缺失或全部解析失败，则输出所有 metric，整块留空。
        metric_keys = dataset_present_metric_keys[dataset]
        if metric_keys:
            metric_keys = sorted(metric_keys, key=lambda k: METRIC_KEY_ORDER[k])
        else:
            metric_keys = [spec["key"] for spec in METRIC_SPECS]

        for key in metric_keys:
            spec = METRIC_KEY_TO_SPEC[key]

            for exp_name in DST_EXP_NAMES:
                for step in STEPS:
                    cached = file_cache[(dataset, exp_name, step)]
                    metrics = cached["metrics"]

                    records.append({
                        "domain": get_domain(dataset),
                        "setting_dir": setting_dir,
                        "setting": setting,
                        "dataset": dataset,
                        "dataset_display": DATASET_DISPLAY.get(dataset, dataset),
                        "task": spec["task"],
                        "metric": spec["metric"],
                        "stat": spec["stat"],
                        "metric_key": key,
                        "exp": exp_name,
                        "step": step,
                        "step_str": step_to_str(step),
                        "value": metrics.get(key, np.nan),
                        "xlsx": cached["xlsx"] if cached["exists"] else "",
                    })

    raw_df = pd.DataFrame(records)
    selected_df = pd.DataFrame(selected_file_records)

    print(f"[INFO] expected xlsx files : {expected_num}")
    print(f"[INFO] missing xlsx files  : {len(missing_files)}")
    print(f"[INFO] failed xlsx files   : {len(failed_files)}")
    print(f"[INFO] parsed xlsx files   : {selected_df['parsed_ok'].sum() if not selected_df.empty else 0}")

    if missing_files:
        print("[WARN] Missing xlsx files, keep blank:")
        for p in missing_files:
            print(f"  {p}")

    if failed_files:
        print("[WARN] Failed to parse xlsx files, keep blank:")
        for p in failed_files:
            print(f"  {p}")

    return raw_df, selected_df


# ============================================================
# 6. 构建宽表
# ============================================================

def build_summary_df(raw_df: pd.DataFrame) -> pd.DataFrame:
    if raw_df.empty:
        return pd.DataFrame()

    raw_df = raw_df.copy()

    target_cols = [
        f"{exp_name}__step{step}"
        for exp_name in DST_EXP_NAMES
        for step in STEPS
    ]

    raw_df["exp_step"] = raw_df.apply(
        lambda r: f"{r['exp']}__step{int(r['step'])}",
        axis=1,
    )

    meta_cols = [
        "domain",
        "setting",
        "setting_dir",
        "dataset",
        "dataset_display",
        "task",
        "metric",
        "stat",
        "metric_key",
    ]

    # 注意：不要用 pivot_table(dropna=False)，会非常慢
    # 这里先 groupby，再 pivot，只保留真实存在过的组合
    value_df = (
        raw_df
        .groupby(meta_cols + ["exp_step"], as_index=False, dropna=False)
        .agg(value=("value", lambda s: s.iloc[0]))
    )

    table = value_df.pivot(
        index=meta_cols,
        columns="exp_step",
        values="value",
    ).reset_index()

    table.columns.name = None

    for col in target_cols:
        if col not in table.columns:
            table[col] = np.nan

    dataset_rank = {d: i for i, d in enumerate(DATASETS)}

    table["domain_rank"] = table["domain"].map(
        lambda x: DOMAIN_ORDER.get(x, 999)
    )
    table["dataset_rank"] = table["dataset"].map(
        lambda x: dataset_rank.get(x, 999)
    )
    table["metric_rank"] = table["metric_key"].map(
        lambda x: METRIC_KEY_ORDER.get(x, 999)
    )

    table = table.sort_values(
        by=["domain_rank", "dataset_rank", "setting", "metric_rank"],
        ascending=[True, True, True, True],
    )

    fixed_cols = [
        "domain",
        "setting",
        "setting_dir",
        "dataset",
        "dataset_display",
        "task",
        "metric",
        "stat",
        "metric_key",
    ]

    return table[fixed_cols + target_cols]


# ============================================================
# 7. 写 Excel：paper-style summary
# ============================================================

def safe_merge(ws, start_row, start_col, end_row, end_col):
    if end_row > start_row or end_col > start_col:
        ws.merge_cells(
            start_row=start_row,
            start_column=start_col,
            end_row=end_row,
            end_column=end_col,
        )


def write_debug_sheet(ws, df: pd.DataFrame):
    if df.empty:
        ws.cell(1, 1, "empty")
        return

    for c_idx, col in enumerate(df.columns, start=1):
        ws.cell(1, c_idx, col)

    for r_idx, (_, row) in enumerate(df.iterrows(), start=2):
        for c_idx, col in enumerate(df.columns, start=1):
            val = row[col]
            if pd.isna(val):
                val = None
            ws.cell(r_idx, c_idx, val)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c in range(1, ws.max_column + 1):
        ws.cell(1, c).font = Font(bold=True)
        ws.cell(1, c).fill = header_fill
        ws.column_dimensions[get_column_letter(c)].width = 20

    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    ws.freeze_panes = "A2"


def style_summary_sheet(ws, max_row: int, max_col: int):
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    setting_fill = PatternFill("solid", fgColor="E2F0D9")
    exp_fill = PatternFill("solid", fgColor="D9EAF7")
    domain_fill = PatternFill("solid", fgColor="FCE4D6")

    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            if isinstance(cell.value, float):
                cell.number_format = "0.000"

    # 根据 section header / step row 着色
    for r in range(1, max_row + 1):
        a_val = ws.cell(r, 1).value
        b_val = ws.cell(r, 2).value
        c_val = ws.cell(r, 3).value
        d_val = ws.cell(r, 4).value

        # section header row: B 有 setting，C/D 为空
        if b_val is not None and c_val is None and d_val is None:
            for c in range(1, max_col + 1):
                ws.cell(r, c).font = Font(bold=True)
                ws.cell(r, c).fill = setting_fill

            for c in range(6, max_col + 1):
                ws.cell(r, c).fill = exp_fill

            if a_val is not None:
                ws.cell(r, 1).fill = domain_fill

        # step row: 上一行是 section header
        if r > 1:
            prev_b = ws.cell(r - 1, 2).value
            prev_c = ws.cell(r - 1, 3).value
            prev_d = ws.cell(r - 1, 4).value

            if prev_b is not None and prev_c is None and prev_d is None:
                for c in range(6, max_col + 1):
                    ws.cell(r, c).font = Font(bold=True)
                    ws.cell(r, c).fill = exp_fill

    ws.freeze_panes = "F3"

    col_widths = {
        1: 16,
        2: 30,
        3: 14,
        4: 16,
        5: 10,
    }

    for c, w in col_widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    for c in range(6, max_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 13

    for r in range(1, max_row + 1):
        ws.row_dimensions[r].height = 20


def write_paper_style_summary(ws, summary_df: pd.DataFrame):
    """
    输出类似 metric_0519.xlsx 的 summary sheet。

    每个 dataset 一个 block：
      Row 1: Domain / Setting / Exp group names
      Row 2: step names
      Then: Dataset / Task / Metric / Stat / values
    """
    ws.title = "summary"

    value_start_col = 6

    target_cols = [
        f"{exp_name}__step{step}"
        for exp_name in DST_EXP_NAMES
        for step in STEPS
    ]

    max_col = value_start_col + len(target_cols) - 1
    current_row = 1

    domain_start_rows = {}
    prev_domain = None

    grouped = summary_df.groupby(
        ["domain", "setting", "setting_dir", "dataset", "dataset_display"],
        sort=False,
        dropna=False,
    )

    for (domain, setting, setting_dir, dataset, dataset_display), block_df in grouped:
        if domain not in domain_start_rows:
            domain_start_rows[domain] = current_row

        # block header row
        if domain != prev_domain:
            ws.cell(current_row, 1, domain)

        ws.cell(current_row, 2, setting)

        col = value_start_col
        for exp_name in DST_EXP_NAMES:
            exp_start_col = col
            exp_end_col = col + len(STEPS) - 1

            ws.cell(
                current_row,
                exp_start_col,
                EXP_DISPLAY_NAMES.get(exp_name, exp_name),
            )
            safe_merge(ws, current_row, exp_start_col, current_row, exp_end_col)

            col = exp_end_col + 1

        # step row
        current_row += 1

        col = value_start_col
        for exp_name in DST_EXP_NAMES:
            for step in STEPS:
                ws.cell(current_row, col, step_to_str(step))
                col += 1

        # setting merge across header and step row
        safe_merge(ws, current_row - 1, 2, current_row, 2)

        # metric rows
        data_start_row = current_row + 1

        block_df = block_df.sort_values(
            by="metric_key",
            key=lambda s: s.map(lambda k: METRIC_KEY_ORDER.get(k, 999)),
        )

        prev_task = None
        prev_metric_key = None
        task_start_row = None
        metric_start_row = None

        for _, row in block_df.iterrows():
            current_row += 1

            task = row["task"]
            metric = row["metric"]
            stat = row["stat"]
            metric_key = row["metric_key"]

            # dataset column only writes once then merge
            if current_row == data_start_row:
                ws.cell(current_row, 2, dataset_display)

            # task merge
            if task != prev_task:
                if task_start_row is not None:
                    safe_merge(ws, task_start_row, 3, current_row - 1, 3)
                task_start_row = current_row
                ws.cell(current_row, 3, task)

            # metric merge
            if metric_key != prev_metric_key:
                if metric_start_row is not None:
                    safe_merge(ws, metric_start_row, 4, current_row - 1, 4)
                metric_start_row = current_row
                ws.cell(current_row, 4, metric)

            ws.cell(current_row, 5, stat)

            col = value_start_col
            for exp_name in DST_EXP_NAMES:
                for step in STEPS:
                    col_name = f"{exp_name}__step{step}"
                    val = row[col_name] if col_name in row.index else np.nan
                    ws.cell(current_row, col, None if pd.isna(val) else float(val))
                    col += 1

            prev_task = task
            prev_metric_key = metric_key

        if task_start_row is not None:
            safe_merge(ws, task_start_row, 3, current_row, 3)

        if metric_start_row is not None:
            safe_merge(ws, metric_start_row, 4, current_row, 4)

        safe_merge(ws, data_start_row, 2, current_row, 2)

        current_row += 1
        prev_domain = domain

    # merge domain column
    domains_in_order = list(domain_start_rows.keys())
    for i, domain in enumerate(domains_in_order):
        start = domain_start_rows[domain]

        if i + 1 < len(domains_in_order):
            end = domain_start_rows[domains_in_order[i + 1]] - 1
        else:
            end = current_row - 1

        safe_merge(ws, start, 1, end, 1)

    style_summary_sheet(ws, current_row - 1, max_col)


def write_excel(summary_df: pd.DataFrame, raw_df: pd.DataFrame, selected_df: pd.DataFrame, out_path: Path):
    wb = Workbook()

    ws_summary = wb.active
    write_paper_style_summary(ws_summary, summary_df)

    ws_raw = wb.create_sheet("raw_long")
    write_debug_sheet(ws_raw, raw_df)

    ws_files = wb.create_sheet("selected_files")
    write_debug_sheet(ws_files, selected_df)

    wb.save(out_path)


# ============================================================
# 8. Main
# ============================================================

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--root",
        type=str,
        default="outputs",
        help="outputs root directory",
    )
    parser.add_argument(
        "--out",
        type=str,
        default="all_metrics_summary.xlsx",
        help="output xlsx path",
    )
    args = parser.parse_args()

    root = Path(args.root).resolve()
    out_path = Path(args.out).resolve()

    print(f"[INFO] root: {root}")
    print(f"[INFO] out : {out_path}")

    if not root.exists():
        raise FileNotFoundError(f"root does not exist: {root}")

    raw_df, selected_df = collect_records(root)

    if raw_df.empty:
        raise RuntimeError("No metric records collected.")

    summary_df = build_summary_df(raw_df)

    if summary_df.empty:
        raise RuntimeError("Summary table is empty.")

    total_values = raw_df["value"].shape[0]
    missing_values = raw_df["value"].isna().sum()

    print(f"[INFO] raw records          : {len(raw_df)}")
    print(f"[INFO] summary rows         : {len(summary_df)}")
    print(f"[INFO] missing metric values: {missing_values}/{total_values}")

    write_excel(
        summary_df=summary_df,
        raw_df=raw_df,
        selected_df=selected_df,
        out_path=out_path,
    )

    print(f"[DONE] saved to: {out_path}")


if __name__ == "__main__":
    main()